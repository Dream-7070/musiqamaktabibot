# -*- coding: utf-8 -*-
# ==========================
# services/schedule_template.py
# DARS JADVALI - EXCEL SHABLONLARI
# ==========================
#
# Ikki tur:
#
#   yakka  - har qator bitta dars: o'quvchi, sinf, fan, kun,
#            vaqt, xona.
#   guruh  - bo'sh qator bilan ajratilgan har blok bitta guruh
#            darsi: blokdagi o'quvchilar - guruh a'zolari, fan/kun/
#            soat/xona blokning birinchi qatoriga yoziladi.
#
# Eski shablonda kun ustundan, fan blok sarlavhasidan olinardi,
# xona esa umuman yozilmasdi - bot taxmin qilardi yoki so'rardi.
#
# Fan, kun, sinf va xona ro'yxatdan tanlanadi (Excel ochiladigan
# ro'yxati) - qo'lda yozilganda imlo xatosi chiqardi.
# ==========================


from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


DAYS = ["Dushanba", "Seshanba", "Chorshanba", "Payshanba", "Juma", "Shanba"]

CLASSES = ["1", "2", "3", "4", "5", "6", "7"]

SLOTS = 3

HEADERS = [
    ("№", 5),
    ("O'quvchining F.I.SH", 30),
    ("Sinfi", 7),
    ("Fan", 30),
]
for i in range(1, SLOTS + 1):
    HEADERS.extend([
        (f"{i}-kun", 13),
        (f"{i}-soat", 14),
        (f"{i}-xona", 9),
    ])

HEADER_ROW = 4

DATA_ROWS = 300


YAKKA_YORIQNOMA = [
    "QANDAY TO'LDIRILADI",
    "",
    "1. Har qator - BITTA o'quvchining BITTA fani. Haftada 2-3 marta bo'lsa - 2-kun/2-soat/2-xona va 3-kun... ustunlarini to'ldiring, qatorni takrorlamang.",
    "   Bir o'quvchining ikki xil fani bo'lsa - ikki qator",
    "",
    "2. Fan, Kun, Sinfi va Xona - ro'yxatdan tanlanadi (katakchani bosing).",
    "",
    "3. Dars soati - boshlanish va tugash vaqti: 9:40-10:25",
    "",
    "4. Ism-familiya botdagi o'quvchilar ro'yxatidagi bilan bir xil bo'lsin.",
    "   Botda yo'q o'quvchining darsi yuklanmaydi - avval uni qo'shing.",
    "",
    "5. Jo'rnavozlik darslarini yozmang - ular mutaxassislik",
    "   o'qituvchisining jadvaliga bog'lanadi.",
    "",
    "NAMUNA:",
    "",
]

YAKKA_NAMUNA = [
    (1, "Turg'inboyev Kamron", "5", "Mutaxassislik", "Dushanba", "9:40-10:50", "2/8", "Chorshanba", "8:50-10:00", "2/8", "", "", ""),
    (2, "Alisherov Zafar", "3", "Notani varaqdan o'qish", "Chorshanba", "10:05-10:50", "2/8", "", "", "", "", "", ""),
]

GURUH_YORIQNOMA = [
    "QANDAY TO'LDIRILADI",
    "",
    "1. Bir blok - BITTA guruh darsi.",
    "   Blokdagi o'quvchilar - shu guruhning a'zolari.",
    "",
    "2. Bloklar orasida BITTA BO'SH QATOR qoldiring.",
    "",
    "3. Fan, Kun, Dars soati va Xona - blokning BIRINCHI qatoriga yoziladi",
    "   (katakchalarni birlashtirsangiz ham bo'ladi).",
    "",
    "4. Guruh haftada 2-3 marta bo'lsa - 2-kun/2-soat/2-xona ustunlarini to'ldiring.",
    "",
    "5. Ism-familiya botdagi o'quvchilar ro'yxatidagi bilan bir xil bo'lsin.",
    "   Botda yo'q bola guruhga qo'shilmaydi - qolganlari bilan guruh qoladi.",
    "",
    "NAMUNA:",
    "",
]

GURUH_NAMUNA = [
    (1, "Turg'inboyev Kamron", "5", "Solfedjio", "Dushanba", "13:00-13:45", "2/8", "Payshanba", "13:00-13:45", "2/8", "", "", ""),
    (2, "Alisherov Zafar", "3", "", "", "", "", "", "", "", "", "", ""),
    (3, "Mansurov Abbosxo'ja", "3", "", "", "", "", "", "", "", "", "", ""),
]


def build_template(path, rooms, subjects, teacher="", guruh=False):
    """Shablonni `path` ga yozadi. rooms, subjects - ro'yxat."""

    wb = Workbook()

    ws = wb.active
    ws.title = "Jadval"

    ws["A1"] = "O'qituvchi: " + (teacher or "______________________")
    ws["A1"].font = Font(bold=True, size=12)

    # Import sarlavha tepasidagi "guruh" so'zidan turini aniqlaydi -
    # shuning uchun bu nom o'zgartirilmasin.

    ws["A2"] = "Guruhli darslar jadvali" if guruh else "Yakka darslar jadvali"
    ws["A2"].font = Font(bold=True, size=14)

    ws["A3"] = "Haftada 2-3 marta bo'lsa - 2-kun, 3-kun ustunlarini to'ldiring"
    ws["A3"].font = Font(italic=True)

    chiziq = Side(style="thin")
    ramka = Border(left=chiziq, right=chiziq, top=chiziq, bottom=chiziq)
    kulrang = PatternFill("solid", fgColor="D9D9D9")
    och_yashil = PatternFill("solid", fgColor="E2EFDA")
    och_sariq = PatternFill("solid", fgColor="FFF2CC")

    for col, (nom, eni) in enumerate(HEADERS, start=1):

        cell = ws.cell(HEADER_ROW, col, nom)
        cell.font = Font(bold=True)
        
        # Color based on slot
        if 8 <= col <= 10:
            cell.fill = och_yashil
        elif 11 <= col <= 13:
            cell.fill = och_sariq
        else:
            cell.fill = kulrang
            
        cell.border = ramka
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)

        ws.column_dimensions[cell.column_letter].width = eni

    first = HEADER_ROW + 1
    last = HEADER_ROW + DATA_ROWS

    for row in range(first, last + 1):
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row, col).border = ramka

    ws.freeze_panes = ws.cell(first, 1)

    # --- ro'yxatlar: alohida yashirin varaqda ---
    #
    # Fanlar 90 dan ortiq - Excel'da ro'yxatni to'g'ridan-to'g'ri
    # yozish 255 belgi bilan cheklangan, shuning uchun katakchalar
    # diapazoniga havola qilinadi.

    ls = wb.create_sheet("Ro'yxatlar")

    for i, nom in enumerate(subjects, start=1):
        ls.cell(i, 1, nom)

    for i, nom in enumerate(rooms, start=1):
        ls.cell(i, 2, str(nom))

    ls.sheet_state = "hidden"

    def _royxat(formula, ustun):

        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = "Ro'yxatdan tanlang"
        dv.errorTitle = "Noto'g'ri qiymat"
        ws.add_data_validation(dv)
        dv.add(ustun + str(first) + ":" + ustun + str(last))

    _royxat('"' + ",".join(CLASSES) + '"', "C")
    _royxat("'Ro''yxatlar'!$A$1:$A$" + str(max(len(subjects), 1)), "D")
    
    from openpyxl.utils import get_column_letter
    for i in range(SLOTS):
        col_offset = 5 + i * 3
        day_col = get_column_letter(col_offset)
        room_col = get_column_letter(col_offset + 2)
        _royxat('"' + ",".join(DAYS) + '"', day_col)
        _royxat("'Ro''yxatlar'!$B$1:$B$" + str(max(len(rooms), 1)), room_col)

    # --- yo'riqnoma ---

    yo = wb.create_sheet("Yo'riqnoma")

    matn = GURUH_YORIQNOMA if guruh else YAKKA_YORIQNOMA
    namunalar = GURUH_NAMUNA if guruh else YAKKA_NAMUNA

    for i, qator in enumerate(matn, start=1):
        yo.cell(i, 1, qator)

    yo["A1"].font = Font(bold=True, size=13)

    namuna_boshi = len(matn) + 1

    for col, (nom, eni) in enumerate(HEADERS, start=1):
        cell = yo.cell(namuna_boshi, col, nom)
        cell.font = Font(bold=True)
        cell.fill = kulrang
        yo.column_dimensions[cell.column_letter].width = eni

    for i, qator in enumerate(namunalar, start=namuna_boshi + 1):
        for col, qiymat in enumerate(qator, start=1):
            if qiymat != "":
                yo.cell(i, col, qiymat)

    wb.save(path)

    return path
