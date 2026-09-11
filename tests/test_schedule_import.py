# -*- coding: utf-8 -*-
"""
services/schedule_import.py - Excel dars jadvalini o'qish.

Sinov haqiqiy fayllarga tayanmaydi: shablon openpyxl bilan shu
yerda quriladi, shuning uchun sinov har qanday kompyuterda
ishlaydi. Shablon o'qituvchilarning haqiqiy jadvalidan nusxa:
satrlar - o'quvchilar, ustunlar - 6 kun, oraliqda fan bloklari.

Tekshiriladigan asosiy holatlar:

  1. Jadval tuzilishi satr raqamiga bog'liq emas (siljish)
  2. Kirill va lotin kun nomlari
  3. Fan bloklari va katak ichidagi qisqartmalar
  4. Guruh darsi - birlashtirilgan katak butun guruhga tegishli
  5. Xato yozilgan fan nomi JIMGINA qabul qilinmaydi
  6. Vaqtdagi imlo xatosi ("14;20") tushuniladi
  7. plan_changes: keep / add / move / remove
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

from services import schedule_import as si


ok, bad = [], []


def check(label, cond):
    (ok if cond else bad).append(label)


TMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_tmp")

os.makedirs(TMP, exist_ok=True)


def build(rows, days=None, first_row=1, title="1-chorak dars jadvali"):
    """
    rows - [(A, B, C, D, E, F, G, H, I), ...] ko'rinishidagi satrlar.
    first_row - jadvalni pastga surish (siljigan shablonni taqlid qilish).
    """

    days = days or ["Dushanba", "Seshanba", "Chorshanba",
                    "Payshanba", "Juma", "Shanba"]

    wb = Workbook()

    ws = wb.active

    ws.title = "1-chorak"

    ws.cell(row=first_row, column=1, value=title)

    ws.cell(row=first_row + 1, column=1,
            value="Doyra sinfi o'qituvchisi Kamolov Oybek")

    header = first_row + 2

    ws.cell(row=header, column=1, value="№")
    ws.cell(row=header, column=2, value="O'quvchilarning F.I.O")
    ws.cell(row=header, column=3, value="Sinfi")

    for offset, day in enumerate(days):
        ws.cell(row=header, column=4 + offset, value=day)

    for index, values in enumerate(rows):

        for column, value in enumerate(values, start=1):

            if value not in (None, ""):
                ws.cell(row=header + 1 + index, column=column, value=value)

    handle, path = tempfile.mkstemp(suffix=".xlsx", dir=TMP)

    os.close(handle)

    wb.save(path)

    return path, ws, header


def read(rows, **kwargs):

    path, _, _ = build(rows, **kwargs)

    try:
        return si.read_workbook(path)

    finally:
        os.remove(path)


def lessons_of(result):

    return result["sheet"]["lessons"]


# ==========================
# 1. ODDIY YAKKA JADVAL
# ==========================


YAKKA = [
    ("", "Mutaxassislik", "", "", "", "", "", "", ""),
    (1, "Murodov Otabek", "3", "8:00-8:45", "", "", "8:50-9:35", "", ""),
    (2, "Kenjaboyev Asadbek", "2", "13:50-14:35", "", "", "", "", ""),
    ("", "N.V.O", "", "", "", "", "", "", ""),
    (1, "Murodov Otabek", "3", "8:50-9:15", "", "", "", "", ""),
    ("", "Tanlangan fan", "", "", "", "", "", "", ""),
    (1, "Kenjaboyev Asadbek", "2", "", "13:00-13:25", "", "", "", ""),
]


result = read(YAKKA)

lessons = lessons_of(result)

check("yakka jadval o'qildi (5 dars)", len(lessons) == 5)

check("fan bloki qo'llandi",
      sorted({lesson["subject"] for lesson in lessons})
      == ["Mutaxassislik", "Notani varaqdan o'qish", "Tanlangan fan"])

check("kun ustundan olindi",
      {lesson["day"] for lesson in lessons}
      == {"Dushanba", "Payshanba", "Seshanba"})

check("vaqt daqiqaga aylandi",
      any(lesson["start"] == 480 and lesson["end"] == 525
          for lesson in lessons))

check("sinf ustuni o'qildi",
      all(lesson["class"] in ("2", "3") for lesson in lessons))

check("savol yo'q - hammasi tanildi", not result["questions"])


# ==========================
# 2. JADVAL SILJIGAN VA KIRILLCHA
# ==========================
#
# Haqiqiy fayllarda sarlavha goh 4-satrda, goh 5-satrda; bir
# varaq kirill, boshqasi lotin. Modul satr raqamiga tayanmasligi
# kerak.


shifted = read(
    YAKKA,
    first_row=6,
    days=["Душанба", "Сешанба", "Чоршанба", "Пайшанба", "Жума", "Шанба"]
)

check("siljigan jadval ham o'qildi", len(lessons_of(shifted)) == 5)

check("kirill kun nomlari lotinga o'girildi",
      {lesson["day"] for lesson in lessons_of(shifted)}
      == {"Dushanba", "Payshanba", "Seshanba"})


# ==========================
# 3. VAQTDAGI IMLO XATOSI
# ==========================


typo = read([
    ("", "Mutaxassislik", "", "", "", "", "", "", ""),
    (1, "Murodov Otabek", "3", "14;20-14:45", "", "", "", "", ""),
])

check("nuqtali vergul bilan yozilgan vaqt tushunildi",
      lessons_of(typo)[0]["start"] == 14 * 60 + 20)


# ==========================
# 4. XATO YOZILGAN FAN NOMI
# ==========================
#
# Eng muhim sinov. Ilgari tanilmagan blok sarlavhasi O'QUVCHI
# deb o'qilardi va undan keyingi darslar OLDINGI fanda qolib
# ketardi - haqiqiy faylda 7 ta jo'rnavozlik darsi shu sababli
# "Tanlangan fan" bo'lib turgan edi.


broken = read([
    ("", "Tanlangan fan", "", "", "", "", "", "", ""),
    (1, "Murodov Otabek", "3", "13:00-13:25", "", "", "", "", ""),
    ("", "Jo'rnavoz;il", "", "", "", "", "", "", ""),
    (1, "Qayumov Qobil", "", "", "14:45-15:30", "", "", "", ""),
])

questions = broken["questions"]

check("xato yozilgan fan nomi savolga chiqdi", len(questions) == 1)

check("taxmin to'g'ri", questions and questions[0]["guess"] == "Jo'rnavozlik")

check("savolda kalit bor", questions and questions[0].get("key"))

after = [
    lesson for lesson in lessons_of(broken)
    if lesson["who"] == "Qayumov Qobil"
]

check("xato sarlavhadan keyingi dars oldingi fanda QOLMADI",
      after and after[0]["subject"] == "Jo'rnavozlik")

check("dars savol kaliti bilan belgilandi",
      after and after[0]["alias_key"] == questions[0]["key"])


# ==========================
# 5. UMUMAN TANILMAGAN FAN
# ==========================


unknown = read([
    ("", "Qandaydir narsa", "", "", "", "", "", "", ""),
    (1, "Murodov Otabek", "3", "13:00-13:45", "", "", "", "", ""),
])

check("noma'lum fan savolga chiqdi", len(unknown["questions"]) == 1)

check("noma'lum fanda taxmin qilinmadi",
      unknown["questions"][0]["guess"] is None)

check("noma'lum fanli dars xato ro'yxatiga tushdi",
      any(issue["level"] == "error" for issue in unknown["issues"]))


# ==========================
# 6. O'QUVCHI SATRI SARLAVHA DEB O'QILMASIN
# ==========================


darssiz = read([
    ("", "Mutaxassislik", "", "", "", "", "", "", ""),
    (1, "Murodov Otabek", "3", "8:00-8:45", "", "", "", "", ""),
    (2, "Tursunov Bobur", "", "", "", "", "", "", ""),
])

check("darsi yo'q o'quvchi savol tug'dirmadi", not darssiz["questions"])


# ==========================
# 7. LUG'AT (aliases)
# ==========================


path, _, _ = build([
    ("", "Jo'rnavoz;il", "", "", "", "", "", "", ""),
    (1, "Qayumov Qobil", "", "", "14:45-15:30", "", "", "", ""),
])

try:
    taught = si.read_workbook(path, aliases={"Jo'rnavoz;il": "Jo'rnavozlik"})

finally:
    os.remove(path)

check("lug'atdagi nom so'ralmaydi", not taught["questions"])

check("lug'at fanni to'g'ri qo'ydi",
      lessons_of(taught)[0]["subject"] == "Jo'rnavozlik")


# ==========================
# 8. GURUH JADVALI
# ==========================
#
# Guruh darsi bir marta yoziladi va birlashtirilgan katak butun
# guruh satrlarini qamraydi.


path, ws, header = build([
    ("", "1-guruh", "", "", "", "", "", "", ""),
    (1, "O'ktamova Umida", "3", "Mut. 13:00-15:15", "", "", "", "", ""),
    (2, "Saidova Muyassar", "3", "", "", "", "", "", ""),
    (3, "Baratova Setora", "2", "", "", "", "", "", ""),
])

wb = load_workbook(path)

sheet = wb["1-chorak"]

first = header + 2

sheet.merge_cells(start_row=first, start_column=4,
                  end_row=first + 2, end_column=4)

wb.save(path)

try:
    group_result = si.read_workbook(path)

finally:
    os.remove(path)

group_lessons = lessons_of(group_result)

check("guruh darsi BIR marta o'qildi", len(group_lessons) == 1)

check("guruh nomi darsga yozildi",
      group_lessons and group_lessons[0]["who"] == "1-guruh")

check("guruh darsi belgilandi",
      group_lessons and group_lessons[0]["is_group"])

check("katakdagi qisqartma fanga o'girildi",
      group_lessons and group_lessons[0]["subject"] == "Mutaxassislik")

check("guruh darsiga sinf qo'yilmadi",
      group_lessons and group_lessons[0]["class"] == "")

members = group_result["sheet"]["groups"].get("1-guruh") or []

check("guruh tarkibi yig'ildi (3 o'quvchi)", len(members) == 3)


# ==========================
# 9. TEKSHIRUV
# ==========================


overlap = read([
    ("", "Mutaxassislik", "", "", "", "", "", "", ""),
    (1, "Murodov Otabek", "3", "8:00-8:45", "", "", "", "", ""),
    (2, "Kenjaboyev Asadbek", "2", "8:30-9:15", "", "", "", "", ""),
])

check("ustma-ust tushgan dars topildi",
      any("ustma-ust" in issue["text"] for issue in overlap["issues"]))


# 80 daqiqa - eng yaqin akademik qiymat 90 (1,5 soat = 65,
# 2 soat = 90). 5 daqiqagacha farq kechiriladi, chunki
# o'qituvchilar vaqtni 5 daqiqalik to'rga yaxlitlaydi.

odd = read([
    ("", "Mutaxassislik", "", "", "", "", "", "", ""),
    (1, "Murodov Otabek", "3", "8:00-9:20", "", "", "", "", ""),
])

check("akademik soatga mos kelmagan davomiylik topildi",
      any("Akademik soatga" in issue["text"] for issue in odd["issues"]))


# ==========================
# 10. NIMA O'ZGARADI
# ==========================


def lesson(day="Dushanba", subject="Mutaxassislik", cls="3",
           who="Murodov Otabek", start=480, room="2/14"):

    return {
        "day": day, "subject": subject, "class": cls, "who": who,
        "start": start, "end": start + 45, "minutes": 45, "room": room,
    }


base = lesson()

plan = si.plan_changes([base], [base])

check("o'zgarmagan dars keep ga tushdi",
      len(plan["keep"]) == 1 and not plan["add"]
      and not plan["move"] and not plan["remove"])


plan = si.plan_changes([base], [lesson(start=510)])

check("vaqti o'zgargan dars move ga tushdi",
      len(plan["move"]) == 1 and not plan["add"] and not plan["remove"])


plan = si.plan_changes([base], [lesson(room="1/23")])

check("xonasi o'zgargan dars move ga tushdi", len(plan["move"]) == 1)


plan = si.plan_changes([base], [base, lesson(day="Juma")])

check("yangi dars add ga tushdi",
      len(plan["add"]) == 1 and len(plan["keep"]) == 1)


plan = si.plan_changes([base, lesson(day="Juma")], [base])

check("faylda yo'q dars remove ga tushdi",
      len(plan["remove"]) == 1 and plan["remove"][0]["day"] == "Juma")


lines = si.describe_changes(
    si.plan_changes([base], [lesson(start=510)])
)

check("o'zgarishlar matni tuzildi",
      any("8:00" in line and "8:30" in line for line in lines))


# ==========================
# 11. FAQAT BIRINCHI VARAQ
# ==========================
#
# O'qituvchi bitta chorak jadvalini bitta faylda yuboradi.
# Faylda eski varaqlar qolib ketgan bo'lsa ham, ular
# e'tiborga olinmasligi kerak.


path, _, _ = build([
    ("", "Mutaxassislik", "", "", "", "", "", "", ""),
    (1, "Murodov Otabek", "3", "8:00-8:45", "", "", "", "", ""),
])

extra = load_workbook(path)

second = extra.create_sheet("Eski")

second["A1"] = "1-chorak dars jadvali"
second["A3"] = "№"
second["B3"] = "O'quvchilarning F.I.O"
second["C3"] = "Sinfi"

for offset, day in enumerate(["Dushanba", "Seshanba", "Chorshanba",
                              "Payshanba", "Juma", "Shanba"]):
    second.cell(row=3, column=4 + offset, value=day)

second["B4"] = "Eskiyev Eski"
second["C4"] = "5"
second["D4"] = "10:00-10:45"

extra.save(path)

try:
    one = si.read_workbook(path)

finally:
    os.remove(path)

check("faqat birinchi varaq o'qildi", len(lessons_of(one)) == 1)

check("ikkinchi varaqdagi dars olinmadi",
      all(item["who"] != "Eskiyev Eski" for item in lessons_of(one)))


# ==========================
# 12. YORDAMCHILAR
# ==========================


check("parse_clock ishlaydi", si.parse_clock("08:05") == 485)

check("parse_clock tushunmasa None", si.parse_clock("nimadir") is None)

check("format_time ishlaydi", si.format_time(485) == "8:05")

check("known_subjects bo'sh emas", len(si.known_subjects()) > 5)

check("normalize apostrofni birxillashtiradi",
      si.normalize("Jo‘rnavozlik") == si.normalize("Jo'rnavozlik"))

check("looks_like_name ismni ajratadi",
      si.looks_like_name("Murodov Otabek")
      and not si.looks_like_name("Jo'rnavoz;il"))


# ==========================
# NATIJA
# ==========================

print()
for line in ok:
    print("  OK   " + line)

for line in bad:
    print("  XATO " + line)

print()
print(str(len(ok)) + " ta o'tdi, " + str(len(bad)) + " ta xato")

sys.exit(1 if bad else 0)
